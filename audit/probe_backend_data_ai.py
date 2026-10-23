"""Isolated audit probes. Never opens the project's existing database.

Run: backend\\.venv\\Scripts\\python.exe audit\\probe_backend_data_ai.py
Auth principal is supplied via dependency override; actual scope/role/queries
and request validation are production code. Fresh temp SQLite per execution.
"""
import asyncio
from datetime import date
import io
import json
import logging
import os
from pathlib import Path
import sys
import tempfile
import uuid
logging.disable(logging.CRITICAL)

ROOT = Path(__file__).resolve().parents[1]
TMP = Path(tempfile.mkdtemp(prefix="backend-data-ai-", dir=ROOT / "audit"))
os.environ["DATABASE_URL"] = "sqlite+aiosqlite:///" + (TMP / "probe.db").as_posix()
os.environ["UPLOAD_DIR"] = str(TMP / "uploads")
os.environ["AI_ENABLED"] = "false"
sys.path.insert(0, str(ROOT / "backend"))

import httpx
from sqlalchemy import select, text, func, delete
from app.api.deps import get_current_user
from app.main import app
from app.db.session import AsyncSessionLocal, async_engine, init_models
from app.models.schools import School
from app.models.staff import Staff
from app.models.academic import SchoolYear, GradeLevel, Section, Subject, TeacherAssignment
from app.models.students import Student, StudentEnrollment, StudentPace, DataQualityLog, AIRiskPrediction
from app.services import grading, ml_client

OBS = []
actor = None

async def current_user():
    return actor

app.dependency_overrides[get_current_user] = current_user

def note(label, **kw):
    row = {"probe": label, **kw}
    OBS.append(row)
    print(json.dumps(row, default=str, ensure_ascii=True))

async def main():
    global actor
    await init_models()
    async with AsyncSessionLocal() as db:
        school = School(name="Audit school", status="approved")
        other = School(name="Other school", status="approved")
        db.add_all([school, other]); await db.flush()
        admin = Staff(school_id=school.school_id, email="admin@audit.example", password_hash="unused", first_name="Admin", last_name="Audit", contact_number="0", role="admin", is_approved=True, account_status="approved")
        teacher = Staff(school_id=school.school_id, email="teacher@audit.example", password_hash="unused", first_name="Math", last_name="Teacher", contact_number="0", role="teacher", is_approved=True, account_status="approved")
        db.add_all([admin, teacher]); await db.flush()
        year = SchoolYear(school_id=school.school_id, year="2026-2027", start_date=date(2026,6,1), end_date=date(2027,5,31), is_current=True)
        level = GradeLevel(school_id=school.school_id, level="7", name="Grade 7")
        db.add_all([year,level]); await db.flush()
        own = Section(school_id=school.school_id, grade_level_id=level.grade_level_id, section_code="OWN", name="Own", capacity=1)
        unassigned = Section(school_id=school.school_id, grade_level_id=level.grade_level_id, section_code="OTHER", name="Unassigned", capacity=40)
        foreign_level = GradeLevel(school_id=other.school_id, level="7", name="Grade 7")
        db.add_all([own,unassigned,foreign_level]); await db.flush()
        foreign_sec = Section(school_id=other.school_id, grade_level_id=foreign_level.grade_level_id, section_code="FOREIGN", name="Foreign")
        math = Subject(school_id=school.school_id,grade_level_id=level.grade_level_id,subject_name="Mathematics",subject_code="MATH")
        science = Subject(school_id=school.school_id,grade_level_id=level.grade_level_id,subject_name="Science",subject_code="SCI")
        db.add_all([math,science,foreign_sec]); await db.flush()
        db.add(TeacherAssignment(school_id=school.school_id,teacher_id=teacher.id,section_id=own.section_id,subject_id=math.subject_id,is_adviser=False))
        own_student = Student(school_id=school.school_id,first_name="Owned",last_name="Learner")
        victim = Student(school_id=school.school_id,first_name="Private",last_name="Learner")
        db.add_all([own_student,victim]); await db.flush()
        en = StudentEnrollment(school_id=school.school_id,student_id=own_student.student_id,grade_level_id=level.grade_level_id,section_id=own.section_id,school_year_id=year.school_year_id,is_active=True)
        victim_en = StudentEnrollment(school_id=school.school_id,student_id=victim.student_id,grade_level_id=level.grade_level_id,section_id=unassigned.section_id,school_year_id=year.school_year_id,is_active=True)
        db.add_all([en,victim_en]); await db.flush()
        log = DataQualityLog(school_id=school.school_id,student_id=victim.student_id,teacher_id=admin.id,issue_type="sensitive",detail="Private guardian document correction")
        db.add(log); await db.commit()
    actor = teacher
    transport = httpx.ASGITransport(app=app, raise_app_exceptions=False)
    async with httpx.AsyncClient(transport=transport, base_url="http://audit.local") as c:
        async def req(method, path, body=None):
            r=await c.request(method,path,json=body)
            try: data=r.json()
            except ValueError: data=r.text
            return r.status_code,data
        note("positive_scope_private_student", result=await req("GET", f"/api/students/{victim.student_id}"))
        note("unassigned_quality_list", result=await req("GET", "/api/data-quality-logs"))
        note("unassigned_quality_resolve", result=await req("PATCH", f"/api/data-quality-logs/{log.log_id}"))
        note("spoof_quality_author", result=await req("POST", "/api/data-quality-logs", {"student_id":victim.student_id,"teacher_id":str(admin.id),"issue_type":"spoofed","detail":"written by unassigned teacher"}))
        # Subject mutation should be refused because this teacher teaches Math only.
        status,created=await req("POST", "/api/paces", {"student_id":own_student.student_id,"enrollment_id":en.enrollment_id,"subject":"Mathematics"})
        note("create_assigned_pace",status=status,pace_id=created.get("pace_id"))
        note("rename_to_unassigned_subject", result=await req("PATCH", f"/api/paces/{created['pace_id']}", {"subject":"Science","current_pace":5}))
        # Supply an inaccessible learner alongside an accessible enrollment.
        status,mismatch=await req("POST", "/api/paces", {"student_id":victim.student_id,"enrollment_id":en.enrollment_id,"subject":"Mathematics"})
        note("mismatched_student_enrollment",status=status,result=mismatch)
        async with AsyncSessionLocal() as db:
            await db.execute(delete(StudentPace))
            from app.models.students import EarlyWarning
            await db.execute(delete(EarlyWarning))
            await db.commit()
        # Header save and bulk are real application paths, never patched.
        setup={"ww_highest":[100],"pt_highest":[100],"exam_highest":[100,100,100]}
        setup_path=f"/api/class-record/setup?section_id={own.section_id}&subject_id={math.subject_id}&grading_period=1"
        note("normal_header",result=await req("PUT",setup_path,setup))
        note("normal_header_read",result=await req("GET",setup_path))
        async def bulk(grade=100, period=1, **extra):
            return await req("POST", "/api/paces/bulk", {"section_id":own.section_id,"subject_id":math.subject_id,"grading_period":period,"entries":[{"student_id":own_student.student_id,"enrollment_id":en.enrollment_id,"subject":"Mathematics","current_pace":0,"expected_pace":0,"total_paces":0,"ww_scores":[grade],"pt_scores":[grade],"exam_scores":[grade,grade,grade],**extra}]})
        note("normal_bulk",result=await bulk())
        status,rows=await req("GET", f"/api/paces?enrollment_id={en.enrollment_id}")
        record=next(x for x in rows if x["subject"]=="Mathematics" and x["grading_period"]==1)
        note("normal_saved_grade", grade=record["transmuted_grade"],record_complete=record["record_complete"])
        async with AsyncSessionLocal() as db:
            sch=await db.get(School,school.school_id);sch.grading_periods=4;await db.commit()
        _,four_list=await req("GET","/api/students")
        _,four_read=await req("GET",f"/api/students/{own_student.student_id}")
        note("four_period_school_required_average",list_needs_to_pass=next(x for x in four_list if x["student_id"]==own_student.student_id)["needs_to_pass"],detail_needs_to_pass=four_read["needs_to_pass"],detail_required_in_weakest=four_read["required_in_weakest"],expected_needs_to_pass=round((75*4-100)/3,2),expected_required_at_80=round((80*4-100)/3,2))
        async with AsyncSessionLocal() as db:
            sch=await db.get(School,school.school_id);sch.grading_periods=3;await db.commit()
        # Partial updates must preserve exactly the same read scope as GET.
        async with AsyncSessionLocal() as db:
            private_science=StudentPace(school_id=school.school_id,student_id=own_student.student_id,enrollment_id=en.enrollment_id,subject="Science",grading_period=1,transmuted_grade=60,record_complete=True)
            db.add(private_science);await db.commit();private_id=private_science.pace_id
        note("subject_scope_get",result=await req("GET",f"/api/students/{own_student.student_id}"))
        note("subject_scope_patch",result=await req("PATCH",f"/api/students/{own_student.student_id}",{"address":"Updated address"}))
        async with AsyncSessionLocal() as db:
            await db.execute(delete(StudentPace).where(StudentPace.pace_id==private_id));await db.commit()
        note("healthy_teacher_reads",results={p: (await req("GET",p))[0] for p in ["/api/students","/api/enrollments","/api/paces","/api/school-years","/api/school-years/current","/api/grade-levels","/api/sections","/api/subjects","/api/teacher-assignments","/api/my-classes","/api/early-warnings","/api/early-warnings/summary","/api/ai/predictions","/api/ai/summary",f"/api/students/{own_student.student_id}/profile"]})
        note("backend_ai_disabled_status",result=await req("GET","/api/ai/status?wake=false"))
        note("backend_ai_disabled_model_card",result=await req("GET","/api/ai/model-card"))
        st,predictions=await req("GET","/api/ai/predictions")
        note("backend_ai_prediction_detail",result=await req("GET",f"/api/ai/predictions/{predictions[0]['prediction_id']}"))
        note("backend_ai_rescore",result=await req("POST","/api/ai/rescore",{}))
        note("teacher_admin_dashboard_denied",result=await req("GET","/api/dashboard/admin"))
        async with AsyncSessionLocal() as db:
            previous=SchoolYear(school_id=school.school_id,year="2025-2026",start_date=date(2025,6,1),end_date=date(2026,5,31),is_current=False)
            db.add(previous);await db.flush()
            historical=StudentEnrollment(school_id=school.school_id,student_id=own_student.student_id,grade_level_id=level.grade_level_id,section_id=own.section_id,school_year_id=previous.school_year_id,is_active=False)
            db.add(historical);await db.flush()
            historic_pace=StudentPace(school_id=school.school_id,student_id=own_student.student_id,enrollment_id=historical.enrollment_id,subject="Mathematics",grading_period=1,transmuted_grade=60,record_complete=True)
            db.add(historic_pace);await db.commit()
        note("historical_grades_current_summary",result=await req("GET",f"/api/students/{own_student.student_id}"))
        async with AsyncSessionLocal() as db:
            await db.execute(delete(StudentPace).where(StudentPace.pace_id==historic_pace.pace_id));await db.execute(delete(StudentEnrollment).where(StudentEnrollment.enrollment_id==historical.enrollment_id));await db.commit()
        note("patch_scores_uses_header",result=await req("PATCH", f"/api/paces/{record['pace_id']}", {"ww_scores":[90]}))
        note("restore_bulk",result=await bulk())
        note("warning_before_header_change", result=await req("GET", "/api/early-warnings"))
        note("lower_grade_via_header",result=await req("PUT",setup_path,{"ww_highest":[200],"pt_highest":[200],"exam_highest":[200,200,200]}))
        note("warnings_after_header_change",result=await req("GET","/api/early-warnings"))
        note("student_after_header_change",result=await req("GET",f"/api/students/{own_student.student_id}"))
        note("dashboard_after_header_change",result=await req("GET","/api/dashboard"))
        # With missing examination values, code claims to withhold a final grade.
        note("incomplete_exam",result=await bulk(100,exam_scores=[100,None,None]))
        note("incomplete_record_read",result=await req("GET",f"/api/paces?enrollment_id={en.enrollment_id}"))
        note("reset_header_before_concurrency",result=await req("PUT",setup_path,setup))
        # Independent async request sessions, each updating the same grade record.
        note("concurrent_same_record_saves",results=await asyncio.gather(*(bulk(g) for g in range(85,90))))
        note("concurrent_same_record_state",result=await req("GET",f"/api/paces?enrollment_id={en.enrollment_id}"))
        note("concurrent_new_record_creation",results=await asyncio.gather(*(bulk(85,period=2) for _ in range(5))))
        # Observe all retry attempts through the real save path, replacing only
        # network responses and delays; no real long-running requests are made.
        from unittest.mock import patch
        from app.core.config import settings
        retries=[];pauses=[]
        class FailedAIClient:
            def __init__(self,**kwargs): retries.append(kwargs)
            async def __aenter__(self): return self
            async def __aexit__(self,*args): pass
            async def request(self,method,url,**kwargs):
                return httpx.Response(503,request=httpx.Request(method,url))
        async def fake_sleep(seconds): pauses.append(seconds)
        settings.AI_ENABLED=True
        with patch.object(ml_client.httpx,"AsyncClient",FailedAIClient), patch.object(ml_client.asyncio,"sleep",fake_sleep):
            note("ordinary_save_ai_outage",result=await bulk(85),configured_clients=retries,retry_pauses=pauses)
        settings.AI_ENABLED=False
        async with AsyncSessionLocal() as db:
            gap_student=Student(school_id=school.school_id,first_name="Gap",last_name="Learner")
            db.add(gap_student);await db.flush()
            gap_en=StudentEnrollment(school_id=school.school_id,student_id=gap_student.student_id,grade_level_id=level.grade_level_id,section_id=own.section_id,school_year_id=year.school_year_id,is_active=True)
            db.add(gap_en);await db.commit()
        await req("PUT",f"/api/class-record/setup?section_id={own.section_id}&subject_id={math.subject_id}&grading_period=3",setup)
        for g in (100,0):
            await req("POST","/api/paces/bulk",{"section_id":own.section_id,"subject_id":math.subject_id,"grading_period":3,"entries":[{"student_id":gap_student.student_id,"enrollment_id":gap_en.enrollment_id,"subject":"Mathematics","current_pace":0,"expected_pace":0,"total_paces":0,"ww_scores":[g],"pt_scores":[g],"exam_scores":[g,g,g]}]})
        note("gap_terms_student",result=await req("GET",f"/api/students/{gap_student.student_id}"))
        st,warning_rows=await req("GET","/api/early-warnings")
        note("gap_terms_warning",status=st,rows=[w for w in warning_rows if w["student_id"]==gap_student.student_id])
        gap_warning=next(w for w in warning_rows if w["student_id"]==gap_student.student_id)
        note("healthy_warning_intervention_patch",result=await req("PATCH",f"/api/early-warnings/{gap_warning['warning_id']}",{"intervention":"Synthetic audit intervention"}))
        # Malformed but ordinary JSON; checks should not result in persistent 500s.
        note("accept_invalid_setup",result=await req("PUT",setup_path,{"ww_highest":[-10],"pt_highest":[100],"exam_highest":[100],"exam_subweights":[-100,100,100]}))
        note("invalid_period_zero",result=await bulk(50,period=0))
        note("invalid_period_negative_accepted",result=await bulk(50,period=-1))
        note("invalid_period_negative",result=await bulk(50,period=-10))
        note("invalid_period_seven",result=await bulk(50,period=7))
        actor=admin
        note("healthy_admin_dashboard",result=await req("GET","/api/dashboard/admin"))
        _,gap_records=await req("GET",f"/api/paces?student_id={gap_student.student_id}")
        note("admin_delete_pace",result=await req("DELETE",f"/api/paces/{gap_records[0]['pace_id']}"))
        _,warnings_after_delete=await req("GET","/api/early-warnings")
        note("warnings_after_pace_delete",remaining_for_deleted_pace_student=[w for w in warnings_after_delete if w["student_id"]==gap_student.student_id])
        note("subject_rename",result=await req("PATCH",f"/api/subjects/{math.subject_id}",{"subject_name":"Renamed Mathematics"}))
        actor=teacher
        note("teacher_paces_after_subject_rename",result=await req("GET",f"/api/paces?enrollment_id={en.enrollment_id}"))
        actor=admin
        await req("PATCH",f"/api/subjects/{math.subject_id}",{"subject_name":"Mathematics"})
        # Academic CRUD through every remaining write route, on disposable rows.
        academic_results=[]
        st,yr=await req("POST","/api/school-years",{"year":"2028-2029"});academic_results.append(["year_create",st])
        academic_results.append(["year_patch",(await req("PATCH",f"/api/school-years/{yr['school_year_id']}",{"year":"2029-2030"}))[0]])
        academic_results.append(["year_delete",(await req("DELETE",f"/api/school-years/{yr['school_year_id']}"))[0]])
        st,lev=await req("POST","/api/grade-levels",{"level":"8","name":"Grade 8"});academic_results.append(["level_create",st])
        academic_results.append(["level_patch",(await req("PATCH",f"/api/grade-levels/{lev['grade_level_id']}",{"name":"Grade Eight"}))[0]])
        st,sec=await req("POST","/api/sections",{"grade_level_id":lev["grade_level_id"],"section_code":"TEMP","name":"Temporary"});academic_results.append(["section_create",st])
        academic_results.append(["section_patch",(await req("PATCH",f"/api/sections/{sec['section_id']}",{"name":"Temporary revised"}))[0]])
        st,sub=await req("POST","/api/subjects",{"grade_level_id":lev["grade_level_id"],"subject_code":"TEMP","subject_name":"Disposable"});academic_results.append(["subject_create",st])
        st,ass=await req("POST","/api/teacher-assignments",{"teacher_id":str(teacher.id),"section_id":sec["section_id"],"subject_id":sub["subject_id"]});academic_results.append(["assignment_create",st])
        academic_results.append(["assignment_delete",(await req("DELETE",f"/api/teacher-assignments/{ass['assignment_id']}"))[0]])
        academic_results.append(["subject_delete",(await req("DELETE",f"/api/subjects/{sub['subject_id']}"))[0]])
        academic_results.append(["section_delete",(await req("DELETE",f"/api/sections/{sec['section_id']}"))[0]])
        academic_results.append(["level_delete",(await req("DELETE",f"/api/grade-levels/{lev['grade_level_id']}"))[0]])
        note("academic_crud",results=academic_results)
        st,tmp_student=await req("POST","/api/students",{"first_name":"Disposable","last_name":"Learner"})
        st,tmp_enrol=await req("POST","/api/enrollments",{"student_id":tmp_student["student_id"],"grade_level_id":level.grade_level_id,"section_id":unassigned.section_id,"school_year_id":year.school_year_id})
        note("enrollment_create_healthy",status=st)
        note("enrollment_patch_into_full_section",result=await req("PATCH",f"/api/enrollments/{tmp_enrol['enrollment_id']}",{"section_id":own.section_id}))
        note("enrollment_delete_healthy",result=await req("DELETE",f"/api/enrollments/{tmp_enrol['enrollment_id']}"))
        note("ungraded_student_delete_healthy",result=await req("DELETE",f"/api/students/{tmp_student['student_id']}"))
        note("crossschool_header",result=await req("PUT",f"/api/class-record/setup?section_id={foreign_sec.section_id}&subject_id={math.subject_id}&grading_period=1",setup))
        note("enroll_above_capacity_via_student",result=await req("POST","/api/students",{"first_name":"Over","last_name":"Capacity","grade_level_id":level.grade_level_id,"section_id":own.section_id,"school_year_id":year.school_year_id}))
        note("invalid_student_update",result=await req("PATCH",f"/api/students/{own_student.student_id}",{"gender":"invalid-gender"}))
        note("invalid_student_read",result=await req("GET",f"/api/students/{own_student.student_id}"))
        note("duplicate_section_update",result=await req("PATCH",f"/api/sections/{unassigned.section_id}",{"section_code":"OWN"}))
        from openpyxl import Workbook
        def workbook_bytes(written,term_scores):
            book=Workbook();book.active.title="INPUT DATA";sheet=book.active
            sheet["C24"]="SUBJECT";sheet["E24"]="Mathematics";sheet["K11"]=written
            for term in (1,2,3):
                s=book.create_sheet(f"TERM {term}");s["B15"]="HIGHEST POSSIBLE SCORE"
                for col in (6,14,20,21,22):
                    s.cell(15,col,100)
                    if term in term_scores:s.cell(18,col,term_scores[term])
            data=io.BytesIO();book.save(data);return data.getvalue()
        async def upload(written,scores,save):
            r=await c.post("/api/class-record/import",data={"section_id":own.section_id,"subject_id":math.subject_id,"save":str(save).lower()},files={"file":("audit.xlsx",workbook_bytes(written,scores),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            try:data=r.json()
            except ValueError:data=r.text
            return r.status_code,data
        async with AsyncSessionLocal() as db:
            twins=[Student(school_id=school.school_id,first_name="Alex",last_name="Twin",middle_name=m) for m in ["A.","B."]]
            db.add_all(twins);await db.flush()
            for twin in twins:
                db.add(StudentEnrollment(school_id=school.school_id,student_id=twin.student_id,grade_level_id=level.grade_level_id,section_id=own.section_id,school_year_id=year.school_year_id,is_active=True))
            await db.commit()
        actor=teacher
        note("same_name_import_preview",result=await upload("TWIN, Alex A.",{1:90},False))
        note("same_name_import_save",result=await upload("TWIN, Alex A.",{1:90},True))
        async with AsyncSessionLocal() as db:
            note("same_name_import_actual_owner",expected_student_id=twins[0].student_id,actual_student_ids=list((await db.execute(select(StudentPace.student_id).where(StudentPace.student_id.in_([t.student_id for t in twins])))).scalars()))
        note("import_partial_failure_preview",result=await upload("PARTIAL, New",{1:90,2:1001},False))
        note("import_partial_failure_save",result=await upload("PARTIAL, New",{1:90,2:1001},True))
        async with AsyncSessionLocal() as db:
            partial=(await db.execute(select(Student).where(Student.last_name=="Partial"))).scalar_one_or_none()
            note("import_partial_committed_state",student_persisted=partial is not None,periods=list((await db.execute(select(StudentPace.grading_period).where(StudentPace.student_id==partial.student_id))).scalars()) if partial else [])
        bad=await c.post("/api/class-record/import",data={"section_id":own.section_id,"subject_id":math.subject_id,"save":"false"},files={"file":("fake.xlsx",b"not an Excel file")})
        note("invalid_import_file",status=bad.status_code,body=bad.json())
        async with AsyncSessionLocal() as db:
            doomed=Student(school_id=school.school_id,first_name="Delete",last_name="Me")
            db.add(doomed);await db.flush()
            prediction=AIRiskPrediction(school_id=school.school_id,student_id=doomed.student_id,subject="Science",features='{"q1_term_grade":60}',risk_score=.9)
            db.add(prediction);await db.commit()
        actor=admin
        note("delete_student_with_prediction",result=await req("DELETE",f"/api/students/{doomed.student_id}"))
        async with AsyncSessionLocal() as db:
            note("state_after_student_delete",student_still_exists=await db.get(Student,doomed.student_id) is not None,prediction_count=(await db.execute(select(func.count()).select_from(AIRiskPrediction).where(AIRiskPrediction.student_id==doomed.student_id))).scalar())
        from app.api.routers.students import delete_student
        from fastapi import Request
        async with AsyncSessionLocal() as db:
            try:
                await delete_student(doomed.student_id,Request({"type":"http","headers":[],"client":("127.0.0.1",1)}),admin,school.school_id,db)
            except Exception as exc:
                note("delete_student_exact_error",error=type(exc).__name__,reason=str(getattr(exc,"orig",exc)))
                await db.rollback()
    async with AsyncSessionLocal() as db:
        note("sqlite_fk_pool",foreign_keys=[(await db.execute(text("PRAGMA foreign_keys"))).scalar()])
        note("section_capacity",capacity=own.capacity,enrolled=(await db.execute(select(func.count()).select_from(StudentEnrollment).where(StudentEnrollment.section_id==own.section_id))).scalar())
    incomplete=grading.compute(ww_scores=[100],pt_scores=[100],exam_scores=[100,None,None],ww_highest=[100],pt_highest=[100],exam_highest=[100,100,100])
    from types import SimpleNamespace
    note("incomplete_pure_grading",**incomplete,ai_first_term_ready=ml_client.first_term_ready(SimpleNamespace(grading_period=1,**{**incomplete,"record_complete":incomplete["complete"]})))
    (ROOT / "audit" / "backend-data-ai-evidence.json").write_text(json.dumps(OBS,indent=2,default=str),encoding="utf-8")
    await async_engine.dispose()

if __name__=="__main__":
    asyncio.run(main())
