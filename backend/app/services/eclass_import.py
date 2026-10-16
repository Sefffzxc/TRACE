"""Reads DepEd's 3-Term E-Class Record (Grades 2-10) workbook.

Only what a teacher typed is read: the learners' names on INPUT DATA, and on
each TERM sheet the HIGHEST POSSIBLE SCORE row, the three weights and the raw
scores. Every figure the workbook computes from them (totals, percentages,
initial and term grades, descriptors) is ignored and worked out again by
app/services/grading.py, so an imported record and a typed one are graded
the same way.

The layout is the official template's, cell for cell:

    INPUT DATA   E13 school ID, E14 school name, E15 school year,
                 E23 teacher, E24 subject, E25 grade level, E26 section,
                 K11:K60 boys' names, N11:N60 girls' names
    TERM 1-3     row 15 highest possible scores and weights,
                 rows 18-67 the boys, rows 69-118 the girls (same order as
                 INPUT DATA), WW1-5 in F:J, PT1-3 in N:P, ST1/ST2/TE in T:V
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Optional

# Upload limit. The blank template is about 400 KB.
MAX_BYTES = 5 * 1024 * 1024

INPUT_SHEET = "INPUT DATA"
TERM_SHEETS = {1: "TERM 1", 2: "TERM 2", 3: "TERM 3"}

NAME_ROWS = range(11, 61)          # 50 per list on INPUT DATA
BOYS_COL, GIRLS_COL = 11, 14       # K, N
BOYS_FIRST_ROW, GIRLS_FIRST_ROW = 18, 69

WW_COLS = (6, 7, 8, 9, 10)         # F:J
PT_COLS = (14, 15, 16)             # N:P
EX_COLS = (20, 21, 22)             # T:V  ST1, ST2, TE
HIGHEST_ROW = 15
WEIGHT_CELLS = ("M15", "S15", "AA15")


class NotAnEClassRecord(ValueError):
    """The file is not the 3-Term E-Class Record, or cannot be read."""


@dataclass
class Learner:
    gender: str                    # "Male" | "Female"
    last_name: str
    first_name: str
    middle_name: Optional[str]
    written: str                   # the name as the file wrote it
    slot: int                      # position in its list, from 0


@dataclass
class Term:
    term: int
    ww_highest: list
    pt_highest: list
    exam_highest: list
    weights: tuple                 # (ww, pt, ex), as percentages
    # learner index -> {"ww": [...], "pt": [...], "exam": [...]}
    scores: dict = field(default_factory=dict)


@dataclass
class EClassRecord:
    header: dict
    learners: list
    terms: list
    warnings: list


# ── Helpers ──────────────────────────────────────────────────────────────────
def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _number(value) -> tuple[Optional[float], Optional[str]]:
    """A score cell as a number. Returns (number, problem)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, bool):
        return None, str(value)
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        try:
            number = float(str(value).strip())
        except ValueError:
            # "A", "Ex", "absent"…: no score was earned or recorded.
            return None, _text(value)
    if number < 0:
        return None, _text(value)
    return number, None


def _proper(word: str) -> str:
    """SANTOS -> Santos, DELA CRUZ -> Dela Cruz; mixed case is left alone."""
    if word.isupper() or word.islower():
        return " ".join(part.capitalize() for part in word.split(" "))
    return word


def split_name(written: str) -> tuple[str, str, Optional[str]]:
    """"SANTOS, Juan Miguel D." -> ("Santos", "Juan Miguel", "D.")."""
    if "," in written:
        last, rest = written.split(",", 1)
    else:
        # No comma: take the last word as the surname.
        parts = written.rsplit(" ", 1)
        last, rest = (parts[1], parts[0]) if len(parts) == 2 else (written, "")
    last, rest = last.strip(), rest.strip()
    middle = None
    words = rest.split(" ")
    # A trailing initial ("D." or "D") is the middle initial.
    if len(words) > 1 and re.fullmatch(r"[A-Za-z]\.?", words[-1]):
        middle = words[-1].upper().rstrip(".") + "."
        words = words[:-1]
    first = " ".join(words).strip()
    return _proper(last), _proper(first), middle


def name_key(last: str, first: str) -> str:
    """How two spellings of one name are compared: case, spacing and
    punctuation do not count."""
    return re.sub(r"[^a-z]", "", f"{last}|{first}".lower().replace("ñ", "n"))


def _trimmed(values: list) -> list:
    """Drop the unused items at the end of a row: an item with no highest
    possible score has not been given."""
    out = list(values)
    while out and out[-1] is None:
        out.pop()
    return out


# ── Reading ──────────────────────────────────────────────────────────────────
def read(content: bytes) -> EClassRecord:
    if len(content) > MAX_BYTES:
        raise NotAnEClassRecord("The file is larger than 5 MB.")
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise NotAnEClassRecord(
            "Upload the class record as an .xlsx file. Older .xls files must be "
            "saved as .xlsx in Excel first."
        )
    try:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — any unreadable file is the same answer
        raise NotAnEClassRecord("This file could not be opened as an Excel workbook.") from exc

    try:
        missing = [name for name in (INPUT_SHEET, *TERM_SHEETS.values())
                   if name not in book.sheetnames]
        if missing:
            raise NotAnEClassRecord(
                "This is not the 3-Term E-Class Record (Grades 2-10). "
                f"Missing sheet: {', '.join(missing)}."
            )
        return _read_book(book)
    finally:
        book.close()


def _grid(sheet, max_row: int, max_col: int) -> dict:
    """The sheet's values as {(row, col): value}, read in one pass (read-only
    workbooks are slow to index cell by cell)."""
    cells = {}
    for r, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True),
        start=1,
    ):
        for c, value in enumerate(row, start=1):
            if value is not None:
                cells[(r, c)] = value
    return cells


def _read_book(book) -> EClassRecord:
    warnings: list[str] = []
    info = _grid(book[INPUT_SHEET], 60, 16)

    if _text(info.get((24, 3))).upper() != "SUBJECT":
        raise NotAnEClassRecord(
            "This is not the 3-Term E-Class Record (Grades 2-10): the INPUT DATA "
            "sheet is not laid out as expected."
        )

    header = {
        "school_id": _text(info.get((13, 5))),
        "school_name": _text(info.get((14, 5))),
        "school_year": _text(info.get((15, 5))),
        "teacher": _text(info.get((23, 5))),
        "subject": _text(info.get((24, 5))),
        "grade_level": _text(info.get((25, 5))),
        "section": _text(info.get((26, 5))),
    }

    learners: list[Learner] = []
    # (gender, slot) -> learner index, to find each learner's row on a term
    by_slot: dict[tuple[str, int], int] = {}
    for gender, col in (("Male", BOYS_COL), ("Female", GIRLS_COL)):
        for slot, row in enumerate(NAME_ROWS):
            written = _text(info.get((row, col)))
            if not written:
                continue
            last, first, middle = split_name(written)
            if not last or not first:
                warnings.append(f'Name "{written}" could not be read as LAST, First; skipped.')
                continue
            by_slot[(gender, slot)] = len(learners)
            learners.append(Learner(gender, last, first, middle, written, slot))

    if not learners:
        raise NotAnEClassRecord("No learners' names were found on the INPUT DATA sheet.")

    seen: dict[str, str] = {}
    for learner in learners:
        key = name_key(learner.last_name, learner.first_name)
        if key in seen:
            warnings.append(f'"{learner.written}" appears twice in the list.')
        seen[key] = learner.written

    terms = []
    for number, sheet_name in TERM_SHEETS.items():
        term = _read_term(book[sheet_name], number, by_slot, learners, warnings)
        if term is not None:
            terms.append(term)

    return EClassRecord(header=header, learners=learners, terms=terms, warnings=warnings)


def _read_term(sheet, number, by_slot, learners, warnings) -> Optional[Term]:
    cells = _grid(sheet, GIRLS_FIRST_ROW + 49, 30)
    label = f"Term {number}"

    if _text(cells.get((HIGHEST_ROW, 2))).upper() != "HIGHEST POSSIBLE SCORE":
        raise NotAnEClassRecord(
            f"This is not the 3-Term E-Class Record (Grades 2-10): {sheet.title} "
            "is not laid out as expected."
        )

    def highest(cols):
        values = []
        for col in cols:
            value, problem = _number(cells.get((HIGHEST_ROW, col)))
            if problem:
                warnings.append(f"{label}: highest possible score \"{problem}\" ignored.")
            values.append(value if value else None)
        return values

    ww_hps, pt_hps, ex_hps = highest(WW_COLS), highest(PT_COLS), highest(EX_COLS)

    from openpyxl.utils import coordinate_to_tuple

    weights = []
    for ref, default in zip(WEIGHT_CELLS, (20.0, 50.0, 30.0)):
        value, _ = _number(cells.get(coordinate_to_tuple(ref)))
        if value is None:
            weights.append(default)
        else:
            weights.append(round(value * 100, 4) if value <= 1 else value)
    if abs(sum(weights) - 100) > 0.01:
        warnings.append(f"{label}: the weights do not add up to 100%; DepEd's 20/50/30 used.")
        weights = [20.0, 50.0, 30.0]

    odd_marks: set[str] = set()
    over: list[str] = []
    scores: dict[int, dict] = {}
    for (gender, slot), index in by_slot.items():
        row = (BOYS_FIRST_ROW if gender == "Male" else GIRLS_FIRST_ROW) + slot
        got = {}
        for key, cols, caps in (("ww", WW_COLS, ww_hps), ("pt", PT_COLS, pt_hps),
                                ("exam", EX_COLS, ex_hps)):
            values = []
            for col, cap in zip(cols, caps):
                value, problem = _number(cells.get((row, col)))
                if problem:
                    odd_marks.add(problem)
                if value is not None and cap is None:
                    warnings.append(
                        f"{label}: {learners[index].written} has a score in a column "
                        "with no highest possible score; ignored."
                    )
                    value = None
                if value is not None and cap is not None and value > cap:
                    over.append(learners[index].written)
                values.append(value)
            got[key] = values
        if any(v is not None for values in got.values() for v in values):
            scores[index] = got

    if not scores:
        return None

    if odd_marks:
        shown = ", ".join(f'"{m}"' for m in sorted(odd_marks)[:5])
        warnings.append(f"{label}: cells marked {shown} were read as blank (no score).")
    if over:
        names = ", ".join(sorted(set(over))[:3])
        more = len(set(over)) - 3
        warnings.append(
            f"{label}: scores above the highest possible score for {names}"
            + (f" and {more} more" if more > 0 else "") + "."
        )

    # Keep only the items actually given: WW and PT drop trailing unused
    # columns; the examinations are always three (ST1, ST2, TE).
    ww_len, pt_len = len(_trimmed(ww_hps)), len(_trimmed(pt_hps))
    for got in scores.values():
        got["ww"] = got["ww"][:ww_len]
        got["pt"] = got["pt"][:pt_len]

    return Term(
        term=number,
        ww_highest=ww_hps[:ww_len],
        pt_highest=pt_hps[:pt_len],
        exam_highest=ex_hps,
        weights=tuple(weights),
        scores=scores,
    )
